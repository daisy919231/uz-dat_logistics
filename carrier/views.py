from django.shortcuts import render
from carrier.models import *
from shipper.models import *
from django.shortcuts import render, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST
from django.shortcuts import redirect
from django.contrib import messages
from django.http import HttpResponseForbidden
import pandas as pd
from openpyxl import Workbook
from rest_framework.response import Response
from django.http import HttpResponse
from carrier.models import Order
from rest_framework.generics import ListAPIView
from carrier.serializers import OrderSerializer
from rest_framework.renderers import TemplateHTMLRenderer, JSONRenderer
from django.utils.timezone import is_aware, make_naive
from io import BytesIO


# Create your views here.
def driver_accept_freight(driver, freight_id, truck, bid=False, bid_price=None):
    freight = Freight.objects.get(id=freight_id)
    truck = Truck.objects.get(driver=driver)

    if freight.status != Freight.StatusChoices.searching:
        raise ValueError("Freight already booked or unavailable")

    order = Order.objects.create(
        truck=truck,
        freight=freight,
        status=Order.StatusChoices.not_complete,
        bid=bid,
        bid_price=bid_price if bid else None,
        avto=truck.truck_number
    )

    freight.status = Freight.StatusChoices.booked
    freight.save()

    return order


@login_required
@require_POST
def driver_accept_freight_view(request):
    #Check if user is a driver
    if not hasattr(request.user, 'driver'):
        messages.error(request, "You are not registered as a driver.")
        return HttpResponseForbidden("You are not authorized to perform this action.")

    driver = request.user.driver
    freight_id = request.POST.get('freight_id')
    truck = Truck.objects.get(driver=driver)
    bid = request.POST.get('bid', False)
    bid_price = request.POST.get('bid_price')

    try:
        order = driver_accept_freight(driver, freight_id, truck, bid=bid, bid_price=bid_price)
        messages.success(request, "Freight accepted successfully!")
        return redirect('carrier:order-detail', order_id=order.id)
    except Exception as e:
        messages.error(request, str(e))
        return redirect('shipper:freight-list')


@login_required
def order_detail(request, order_id):
    order = get_object_or_404(Order, id=order_id, truck__driver__user=request.user)
    return render(request, 'carrier/order.html', {'order': order})


@login_required
def driver_orders(request):
    # Get orders for the current driver user
    orders = Order.objects.filter(truck__driver=request.user.driver)
    
    context = {
        'orders': orders,
        'title': 'My Orders'
    }
    return render(request, 'carrier/driver_orders.html', context)

class AllOrdersView(ListAPIView):
    queryset = Order.objects.select_related('truck', 'freight').all()
    serializer_class = OrderSerializer
    authentication_classes = [SessionAuthentication]
    permission_classes = [IsAuthenticated]
    renderer_classes = [JSONRenderer, TemplateHTMLRenderer]
    template_name = 'carrier/all_orders.html'  # Combined template

    def get(self, request, *args, **kwargs):
        orders = self.get_queryset()
        serializer = self.get_serializer(orders, many=True)
        return Response({'orders': serializer.data})
    
def make_naive_safe(dt):
    if dt and is_aware(dt):
        return make_naive(dt)
    return dt

def export_orders_to_excel(request):
    wb = Workbook()  # ✅ use Workbook(), not openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"

    headers = ['ID', 'Status', 'Bid', 'Bid Price', 'Avto', 'Truck', 'Freight Origin', 'Freight Destination', 'Distance', 'Offered Price', 'Created At']
    ws.append(headers)

    for order in Order.objects.select_related('freight', 'truck'):
        created_at = make_naive_safe(getattr(order, 'created_at', None))
        ws.append([
            order.id,
            order.status,
            "Yes" if order.bid else "No",
            str(order.bid_price) if order.bid_price else "",
            order.avto or "",
            str(order.truck),
            order.freight.origin or "",
            order.freight.destination or "",
            order.freight.distance or 0,
            order.freight.offered_price or 0.0,
            created_at,
        ])

    # Save to in-memory file
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=orders.xlsx'
    return response